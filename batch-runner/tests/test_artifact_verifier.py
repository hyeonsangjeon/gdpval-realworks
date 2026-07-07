"""Tests for core/artifact_verifier.py — open valid fixtures, flag bad ones.

Fixtures are generated with the same libraries that validate them, so each test
``importorskip``s its dependency and is robust in a light environment.
"""

import json
import zipfile
from pathlib import Path

import pytest

from core.artifact_verifier import (
    classify_kind,
    verify_artifacts,
    verify_one,
)


def test_classify_kind():
    assert classify_kind(".xlsx") == "spreadsheet"
    assert classify_kind(".PDF") == "pdf"
    assert classify_kind(".mp4") == "video"
    assert classify_kind(".unknownext") == "unknown"


def test_empty_file_is_blocking(tmp_path):
    f = tmp_path / "empty.xlsx"
    f.write_bytes(b"")
    r = verify_one(f, workdir=tmp_path)
    assert r.openable is False
    assert any("empty" in e for e in r.errors)


def test_missing_file(tmp_path):
    r = verify_one(tmp_path / "nope.pdf", workdir=tmp_path)
    assert r.exists is False
    assert r.errors


def test_path_escape_is_blocking(tmp_path):
    outside = tmp_path.parent / "escape.txt"
    outside.write_text("hi")
    try:
        r = verify_one(outside, workdir=tmp_path)
        assert any("escape" in e for e in r.errors)
    finally:
        outside.unlink(missing_ok=True)


def test_valid_xlsx(tmp_path):
    openpyxl = pytest.importorskip("openpyxl")
    f = tmp_path / "a.xlsx"
    wb = openpyxl.Workbook()
    wb.active["A1"] = "hi"
    wb.save(f)
    r = verify_one(f, workdir=tmp_path)
    assert r.openable is True
    assert r.metadata["sheet_count"] == 1
    assert r.sha256


def test_corrupt_xlsx_flagged(tmp_path):
    pytest.importorskip("openpyxl")
    f = tmp_path / "bad.xlsx"
    f.write_bytes(b"this is not a real workbook")
    r = verify_one(f, workdir=tmp_path)
    assert r.openable is False
    assert r.errors


def test_valid_docx(tmp_path):
    docx = pytest.importorskip("docx")
    f = tmp_path / "b.docx"
    d = docx.Document()
    d.add_paragraph("hello world")
    d.save(f)
    r = verify_one(f, workdir=tmp_path)
    assert r.openable is True
    assert r.metadata["paragraph_count"] >= 1


def test_valid_pptx(tmp_path):
    pptx = pytest.importorskip("pptx")
    f = tmp_path / "c.pptx"
    prs = pptx.Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    prs.save(f)
    r = verify_one(f, workdir=tmp_path)
    assert r.openable is True
    assert r.metadata["slide_count"] == 1


def test_valid_pdf(tmp_path):
    fitz = pytest.importorskip("fitz")
    f = tmp_path / "e.pdf"
    doc = fitz.open()
    doc.new_page()
    doc.save(f)
    doc.close()
    r = verify_one(f, workdir=tmp_path)
    assert r.openable is True
    assert r.metadata["page_count"] == 1


def test_valid_png(tmp_path):
    PIL = pytest.importorskip("PIL.Image")
    from PIL import Image
    f = tmp_path / "f.png"
    Image.new("RGB", (12, 8), "blue").save(f)
    r = verify_one(f, workdir=tmp_path)
    assert r.openable is True
    assert r.metadata["width"] == 12 and r.metadata["height"] == 8


def test_valid_json(tmp_path):
    f = tmp_path / "g.json"
    f.write_text(json.dumps({"a": 1}))
    r = verify_one(f, workdir=tmp_path)
    assert r.openable is True


def test_invalid_json_flagged(tmp_path):
    f = tmp_path / "g.json"
    f.write_text("{not valid json,,")
    r = verify_one(f, workdir=tmp_path)
    assert r.openable is False


def test_valid_zip(tmp_path):
    f = tmp_path / "h.zip"
    with zipfile.ZipFile(f, "w") as z:
        z.writestr("inner.txt", "data")
    r = verify_one(f, workdir=tmp_path)
    assert r.openable is True
    assert r.metadata["entries"] == 1


def test_text_file(tmp_path):
    f = tmp_path / "notes.txt"
    f.write_text("some notes\nsecond line")
    r = verify_one(f, workdir=tmp_path)
    assert r.openable is True
    assert r.metadata["line_count"] == 2


def test_verify_artifacts_rollup(tmp_path):
    pytest.importorskip("openpyxl")
    import openpyxl
    good = tmp_path / "good.xlsx"
    wb = openpyxl.Workbook(); wb.active["A1"] = "x"; wb.save(good)
    empty = tmp_path / "empty.docx"
    empty.write_bytes(b"")
    report = verify_artifacts([good, empty], workdir=tmp_path)
    assert report.ok is False
    assert any("empty.docx" in e for e in report.blocking_errors)
    assert len(report.artifacts) == 2
